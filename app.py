
import re
from io import BytesIO

import fitz
from docx import Document
from flask import Flask, render_template, request
from werkzeug.utils import secure_filename
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

ALLOWED_EXTENSIONS = {"pdf", "docx"}

SKILLS = [
    "python", "java", "javascript", "typescript", "c++", "c#",
    "html", "css", "react", "node.js", "express", "flask", "django",
    "sql", "mysql", "postgresql", "mongodb", "firebase",
    "machine learning", "deep learning", "artificial intelligence",
    "data analysis", "data science", "pandas", "numpy", "scikit-learn",
    "tensorflow", "pytorch", "excel", "power bi", "tableau",
    "git", "github", "docker", "kubernetes", "aws", "azure",
    "communication", "leadership", "problem solving", "teamwork",
    "project management", "rest api", "linux"
]

JOBS = [
    {
        "title": "Python Developer",
        "skills": ["python", "flask", "django", "sql", "git",
                   "rest api", "mysql", "mongodb"],
        "description": "Develop Python applications, APIs and backend services."
    },
    {
        "title": "Web Developer",
        "skills": ["html", "css", "javascript", "react", "node.js",
                   "express", "git", "rest api"],
        "description": "Build responsive websites and web applications."
    },
    {
        "title": "Data Analyst",
        "skills": ["python", "sql", "excel", "power bi", "tableau",
                   "pandas", "numpy", "data analysis"],
        "description": "Analyze data and create reports and visualizations."
    },
    {
        "title": "Machine Learning Engineer",
        "skills": ["python", "machine learning", "deep learning",
                   "scikit-learn", "tensorflow", "pytorch", "pandas", "numpy"],
        "description": "Build and evaluate machine learning models."
    },
    {
        "title": "Java Developer",
        "skills": ["java", "sql", "rest api", "git", "mysql", "linux"],
        "description": "Develop Java-based software applications."
    },
    {
        "title": "DevOps Engineer",
        "skills": ["linux", "docker", "kubernetes", "aws", "azure", "git"],
        "description": "Support cloud infrastructure, deployment and automation."
    }
]


def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
    )


def extract_text(uploaded, extension):
    file_bytes = uploaded.read()

    if extension == "pdf":
        with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
            return "\n".join(page.get_text() for page in pdf)

    if extension == "docx":
        document = Document(BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs]
        table_text = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return "\n".join(paragraphs + table_text)

    return ""


def find_skills(text):
    text_lower = text.lower()
    found = []

    for skill in SKILLS:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
        if re.search(pattern, text_lower):
            found.append(skill)

    return sorted(set(found))


def calculate_similarity(resume_text, job_text):
    try:
        vectorizer = TfidfVectorizer(stop_words="english")
        vectors = vectorizer.fit_transform([resume_text, job_text])
        score = cosine_similarity(vectors[0:1], vectors[1:2])[0][0]
        return round(float(score) * 100)
    except ValueError:
        return 0


def match_jobs(resume_text, resume_skills):
    results = []

    for job in JOBS:
        required = job["skills"]
        matched = [s for s in required if s in resume_skills]
        missing = [s for s in required if s not in resume_skills]

        skill_score = round(len(matched) / len(required) * 100)
        job_text = job["title"] + " " + job["description"] + " " + " ".join(required)
        similarity_score = calculate_similarity(resume_text, job_text)

        score = round(skill_score * 0.7 + similarity_score * 0.3)

        results.append({
            "title": job["title"],
            "description": job["description"],
            "score": score,
            "skill_score": skill_score,
            "matched_skills": matched,
            "missing_skills": missing,
            "required_skills": required
        })

    return sorted(results, key=lambda item: item["score"], reverse=True)


def calculate_category_scores(resume_text):
    text = resume_text.lower()
    words = text.split()

    education_keywords = [
        "education", "bachelor", "master", "b.tech", "m.tech",
        "bca", "mca", "degree", "university", "college", "diploma"
    ]
    project_keywords = [
        "project", "projects", "developed", "built",
        "implemented", "github", "application"
    ]
    experience_keywords = [
        "experience", "internship", "intern", "worked",
        "employment", "professional", "responsibilities"
    ]

    def keyword_score(keywords):
        found = sum(1 for keyword in keywords if keyword in text)
        return round(found / len(keywords) * 100)

    content_score = min(100, round(len(words) / 4))
    education_score = keyword_score(education_keywords)
    projects_score = keyword_score(project_keywords)
    experience_score = keyword_score(experience_keywords)

    return content_score, education_score, projects_score, experience_score


@app.route("/")
def index():
    return render_template("index.html", error=None)


@app.route("/analyze", methods=["POST"])
def analyze_resume():
    uploaded = request.files.get("resume")

    if not uploaded or not uploaded.filename:
        return render_template(
            "index.html",
            error="Please choose a PDF or DOCX resume."
        )

    if not allowed_file(uploaded.filename):
        return render_template(
            "index.html",
            error="Only PDF and DOCX files are supported."
        )

    try:
        filename = secure_filename(uploaded.filename)
        extension = filename.rsplit(".", 1)[1].lower()
        resume_text = extract_text(uploaded, extension)

        if not resume_text.strip():
            return render_template(
                "index.html",
                error="No readable text found. Scanned PDFs may need OCR."
            )

        resume_skills = find_skills(resume_text)
        job_results = match_jobs(resume_text, resume_skills)

        overall_score = job_results[0]["score"] if job_results else 0
        word_count = len(resume_text.split())

        content_score, education_score, projects_score, experience_score = (
            calculate_category_scores(resume_text)
        )

        suggestions = []

        if not resume_skills:
            suggestions.append(
                "Add a clearly labelled skills section to your resume."
            )

        if word_count < 100:
            suggestions.append(
                "Your resume text is short. Consider adding projects, "
                "education, experience and achievements."
            )

        if job_results and job_results[0]["missing_skills"]:
            suggestions.append(
                "Consider developing these skills for your top matching role, "
                "if relevant: "
                + ", ".join(job_results[0]["missing_skills"]) + "."
            )

        if experience_score < 20:
            suggestions.append(
                "Make your work experience or internship details clearer, "
                "if you have relevant experience."
            )

        if education_score < 20:
            suggestions.append(
                "Make your education section easy to identify."
            )

        if projects_score < 20:
            suggestions.append(
                "Add a projects section describing your contributions and tools."
            )

        if not suggestions:
            suggestions.append(
                "Tailor your resume to relevant roles and include measurable "
                "achievements where possible."
            )

        return render_template(
            "result.html",
            filename=filename,
            resume_skills=resume_skills,
            job_results=job_results,
            overall_score=overall_score,
            suggestions=suggestions,
            word_count=word_count,
            content_score=content_score,
            education_score=education_score,
            projects_score=projects_score,
            experience_score=experience_score
        )

    except Exception:
        app.logger.exception("Resume analysis failed")
        return render_template(
            "index.html",
            error="Could not analyze this file. Please try another PDF or DOCX."
        )


@app.errorhandler(413)
def file_too_large(_error):
    return "File too large. Please upload a file smaller than 10 MB.", 413


if __name__ == "__main__":
    app.run(debug=True)
